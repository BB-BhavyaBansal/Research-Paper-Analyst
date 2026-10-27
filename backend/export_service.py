import io
from typing import Dict, Any, List
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_background(cell, fill_hex: str):
    """Sets background color of a Word table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def generate_docx_report(analysis: Dict[str, Any], metadata: Dict[str, Any]) -> io.BytesIO:
    """
    Generates a professionally styled Microsoft Word (.docx) document
    containing all 7 core sections of the research paper analysis.
    """
    doc = Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Lavender primary color
    COLOR_PRIMARY = RGBColor(109, 40, 217)   # Deep Lavender / Violet #6D28D9
    COLOR_DARK = RGBColor(30, 27, 75)        # Dark Slate #1E1B4B
    COLOR_MUTED = RGBColor(100, 116, 139)    # Slate 500 #64748B

    # Document Header Badge
    badge_para = doc.add_paragraph()
    badge_run = badge_para.add_run("AI RESEARCH PAPER ANALYSIS REPORT")
    badge_run.font.name = "Arial"
    badge_run.font.size = Pt(9)
    badge_run.font.bold = True
    badge_run.font.color.rgb = COLOR_PRIMARY
    badge_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    badge_para.paragraph_format.space_after = Pt(4)

    # Document Title
    paper_title = analysis.get("paper_title") or metadata.get("title") or metadata.get("filename", "Research Paper Analysis")
    title_para = doc.add_paragraph()
    title_run = title_para.add_run(paper_title)
    title_run.font.name = "Arial"
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = COLOR_DARK
    title_para.paragraph_format.space_after = Pt(8)

    # Metadata Card (Table)
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    authors = ", ".join(analysis.get("authors", [])) if analysis.get("authors") else "Not stated in paper"
    domain = analysis.get("core_domain", "Research Paper")
    file_info = f"{metadata.get('filename', 'Unknown')} ({metadata.get('total_pages', 0)} pages)"
    pub_info = analysis.get("publication_info", "Not specified")

    meta_items = [
        ("Authors", authors),
        ("Core Domain", domain),
        ("File Source", file_info),
        ("Publication Info", pub_info)
    ]

    for idx, (label, val) in enumerate(meta_items):
        row = idx // 2
        col = idx % 2
        cell = table.cell(row, col)
        set_cell_background(cell, "F5F3FF")  # Soft lavender tint
        cell_p = cell.paragraphs[0]
        cell_p.paragraph_format.space_after = Pt(2)
        l_run = cell_p.add_run(f"{label}: ")
        l_run.font.bold = True
        l_run.font.size = Pt(9.5)
        l_run.font.color.rgb = COLOR_PRIMARY
        v_run = cell_p.add_run(str(val))
        v_run.font.size = Pt(9.5)
        v_run.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_section_header(num: int, title: str):
        h_para = doc.add_paragraph()
        h_para.paragraph_format.space_before = Pt(16)
        h_para.paragraph_format.space_after = Pt(6)
        h_run = h_para.add_run(f"{num}. {title}")
        h_run.font.name = "Arial"
        h_run.font.size = Pt(13)
        h_run.font.bold = True
        h_run.font.color.rgb = COLOR_PRIMARY

    # 1. Paper Summary
    add_section_header(1, "Paper Summary")
    sum_p = doc.add_paragraph(analysis.get("paper_summary", "No summary available."))
    sum_p.style.font.name = "Arial"
    sum_p.style.font.size = Pt(10.5)

    # 2. Research Problem and Objectives
    add_section_header(2, "Research Problem and Objectives")
    rpo = analysis.get("research_problem_and_objectives", {})
    if isinstance(rpo, dict):
        if rpo.get("problem_statement"):
            p1 = doc.add_paragraph()
            r1 = p1.add_run("Problem Statement: ")
            r1.font.bold = True
            p1.add_run(rpo.get("problem_statement"))
        
        if rpo.get("motivation"):
            p2 = doc.add_paragraph()
            r2 = p2.add_run("Motivation: ")
            r2.font.bold = True
            p2.add_run(rpo.get("motivation"))

        if rpo.get("objectives"):
            p3 = doc.add_paragraph()
            r3 = p3.add_run("Core Objectives:")
            r3.font.bold = True
            for obj in rpo.get("objectives", []):
                bp = doc.add_paragraph(obj, style='List Bullet')
                bp.paragraph_format.space_after = Pt(2)
    else:
        doc.add_paragraph(str(rpo))

    # 3. Key Contributions and Important Points
    add_section_header(3, "Key Contributions and Important Points")
    contributions = analysis.get("key_contributions", [])
    if isinstance(contributions, list):
        for item in contributions:
            if isinstance(item, dict):
                p = doc.add_paragraph(style='List Bullet')
                t_run = p.add_run(f"{item.get('title', 'Contribution')}: ")
                t_run.font.bold = True
                p.add_run(item.get("description", ""))
                if item.get("page_reference"):
                    ref_run = p.add_run(f" [{item.get('page_reference')}]")
                    ref_run.font.italic = True
                    ref_run.font.color.rgb = COLOR_MUTED
            else:
                doc.add_paragraph(str(item), style='List Bullet')
    else:
        doc.add_paragraph(str(contributions))

    # 4. Research Methodology
    add_section_header(4, "Research Methodology")
    meth = analysis.get("research_methodology", {})
    if isinstance(meth, dict):
        for k, v in [
            ("Approach Overview", meth.get("approach_overview")),
            ("Architectural & Experimental Details", meth.get("architectural_details")),
            ("Baseline Comparisons", meth.get("baseline_comparisons")),
            ("Implementation & Reproducibility", meth.get("implementation_details"))
        ]:
            if v:
                mp = doc.add_paragraph()
                mr = mp.add_run(f"{k}: ")
                mr.font.bold = True
                mp.add_run(v)
    else:
        doc.add_paragraph(str(meth))

    # 5. Results and Findings
    add_section_header(5, "Results and Findings")
    res = analysis.get("results_and_findings", {})
    if isinstance(res, dict):
        if res.get("summary"):
            doc.add_paragraph(res.get("summary"))
        if res.get("metrics_and_benchmarks"):
            doc.add_paragraph().add_run("Key Metrics & Benchmarks:").font.bold = True
            for m in res.get("metrics_and_benchmarks", []):
                doc.add_paragraph(m, style='List Bullet')
        if res.get("significance"):
            sp = doc.add_paragraph()
            sr = sp.add_run("Significance: ")
            sr.font.bold = True
            sp.add_run(res.get("significance"))
    else:
        doc.add_paragraph(str(res))

    # 6. Limitations
    add_section_header(6, "Limitations")
    limits = analysis.get("limitations", [])
    if isinstance(limits, list):
        for lim in limits:
            doc.add_paragraph(lim, style='List Bullet')
    else:
        doc.add_paragraph(str(limits))

    # 7. Future Research Directions
    add_section_header(7, "Future Research Directions")
    fut = analysis.get("future_research_directions", [])
    if isinstance(fut, list):
        for f in fut:
            doc.add_paragraph(f, style='List Bullet')
    else:
        doc.add_paragraph(str(fut))

    # Footer note
    doc.add_paragraph().paragraph_format.space_before = Pt(24)
    footer_p = doc.add_paragraph()
    f_run = footer_p.add_run("Generated by Research Paper Analyst • Powered by Google Gemini AI")
    f_run.font.size = Pt(8.5)
    f_run.font.italic = True
    f_run.font.color.rgb = COLOR_MUTED
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio


def generate_printable_html(analysis: Dict[str, Any], metadata: Dict[str, Any]) -> str:
    """
    Generates a clean, professional, responsive HTML document optimized for
    direct browser viewing and clean PDF printing via Ctrl+P / window.print().
    """
    title = analysis.get("paper_title") or metadata.get("title") or metadata.get("filename", "Research Paper Analysis")
    authors = ", ".join(analysis.get("authors", [])) if analysis.get("authors") else "Not stated in paper"
    domain = analysis.get("core_domain", "Academic Research")
    filename = metadata.get("filename", "")
    pages = metadata.get("total_pages", "")

    # Format Key Contributions
    contribs_html = ""
    for c in analysis.get("key_contributions", []):
        if isinstance(c, dict):
            ref = f'<span class="ref">{c.get("page_reference", "")}</span>' if c.get("page_reference") else ""
            contribs_html += f'<li><strong>{c.get("title", "")}:</strong> {c.get("description", "")} {ref}</li>'
        else:
            contribs_html += f'<li>{c}</li>'

    # Format Limitations
    limits_html = "".join([f"<li>{lim}</li>" for lim in analysis.get("limitations", [])])
    
    # Format Future directions
    future_html = "".join([f"<li>{f}</li>" for f in analysis.get("future_research_directions", [])])

    rpo = analysis.get("research_problem_and_objectives", {})
    objectives_html = ""
    if isinstance(rpo, dict) and rpo.get("objectives"):
        objectives_html = "<ul>" + "".join([f"<li>{o}</li>" for o in rpo.get("objectives", [])]) + "</ul>"

    meth = analysis.get("research_methodology", {})
    res = analysis.get("results_and_findings", {})
    metrics_html = ""
    if isinstance(res, dict) and res.get("metrics_and_benchmarks"):
        metrics_html = "<ul>" + "".join([f"<li>{m}</li>" for m in res.get("metrics_and_benchmarks", [])]) + "</ul>"

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <title>{title} - Research Paper Analysis</title>
  <style>
    @import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700&display=swap');
    
    @page {{
      size: A4;
      margin: 20mm;
    }}
    
    body {{
      font-family: 'Plus Jakarta Sans', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
      color: #1e1b4b;
      background: #ffffff;
      line-height: 1.6;
      margin: 0;
      padding: 30px;
    }}
    
    .header {{
      border-bottom: 2px solid #ede9fe;
      padding-bottom: 20px;
      margin-bottom: 25px;
    }}
    
    .badge {{
      display: inline-block;
      background: #f5f3ff;
      color: #7c3aed;
      font-size: 11px;
      font-weight: 700;
      letter-spacing: 0.08em;
      text-transform: uppercase;
      padding: 4px 10px;
      border-radius: 6px;
      margin-bottom: 12px;
      border: 1px solid #ddd6fe;
    }}
    
    h1 {{
      font-size: 24px;
      font-weight: 700;
      color: #1e1b4b;
      margin: 0 0 16px 0;
      line-height: 1.3;
    }}
    
    .meta-grid {{
      display: grid;
      grid-template-columns: 1fr 1fr;
      gap: 12px;
      background: #faf5ff;
      border: 1px solid #e9d5ff;
      padding: 14px 18px;
      border-radius: 10px;
      font-size: 13px;
    }}
    
    .meta-item strong {{
      color: #6d28d9;
    }}
    
    .section {{
      margin-top: 26px;
      page-break-inside: avoid;
    }}
    
    h2 {{
      font-size: 16px;
      font-weight: 700;
      color: #6d28d9;
      border-left: 4px solid #7c3aed;
      padding-left: 10px;
      margin: 0 0 10px 0;
    }}
    
    p {{
      font-size: 13.5px;
      color: #334155;
      margin: 0 0 10px 0;
    }}
    
    ul {{
      margin: 6px 0 12px 20px;
      padding: 0;
      font-size: 13.5px;
      color: #334155;
    }}
    
    li {{
      margin-bottom: 6px;
    }}
    
    .ref {{
      font-size: 11px;
      color: #8b5cf6;
      background: #f5f3ff;
      padding: 2px 6px;
      border-radius: 4px;
      font-weight: 600;
    }}
    
    .print-bar {{
      position: sticky;
      top: 0;
      background: #ffffff;
      padding: 12px 0;
      border-bottom: 1px solid #e2e8f0;
      margin-bottom: 20px;
      display: flex;
      justify-content: flex-end;
      gap: 10px;
    }}
    
    .btn {{
      background: #7c3aed;
      color: white;
      border: none;
      padding: 8px 16px;
      border-radius: 6px;
      font-weight: 600;
      font-size: 13px;
      cursor: pointer;
    }}
    
    @media print {{
      .print-bar {{ display: none; }}
      body {{ padding: 0; }}
    }}
  </style>
</head>
<body>
  <div class="print-bar">
    <button class="btn" onclick="window.print()">Print / Save as PDF</button>
  </div>

  <div class="header">
    <div class="badge">Research Paper Analysis</div>
    <h1>{title}</h1>
    <div class="meta-grid">
      <div class="meta-item"><strong>Authors:</strong> {authors}</div>
      <div class="meta-item"><strong>Core Domain:</strong> {domain}</div>
      <div class="meta-item"><strong>Document:</strong> {filename} ({pages} pages)</div>
      <div class="meta-item"><strong>Publication:</strong> {analysis.get("publication_info", "Not specified")}</div>
    </div>
  </div>

  <div class="section">
    <h2>1. Paper Summary</h2>
    <p>{analysis.get("paper_summary", "No summary available.")}</p>
  </div>

  <div class="section">
    <h2>2. Research Problem and Objectives</h2>
    <p><strong>Problem Statement:</strong> {rpo.get("problem_statement", "Not stated") if isinstance(rpo, dict) else rpo}</p>
    {f'<p><strong>Motivation:</strong> {rpo.get("motivation")}</p>' if isinstance(rpo, dict) and rpo.get("motivation") else ""}
    {objectives_html}
  </div>

  <div class="section">
    <h2>3. Key Contributions and Important Points</h2>
    <ul>{contribs_html}</ul>
  </div>

  <div class="section">
    <h2>4. Research Methodology</h2>
    <p><strong>Approach Overview:</strong> {meth.get("approach_overview", "Not stated") if isinstance(meth, dict) else meth}</p>
    {f'<p><strong>Architecture & Experiments:</strong> {meth.get("architectural_details")}</p>' if isinstance(meth, dict) and meth.get("architectural_details") else ""}
    {f'<p><strong>Baseline Comparisons:</strong> {meth.get("baseline_comparisons")}</p>' if isinstance(meth, dict) and meth.get("baseline_comparisons") else ""}
  </div>

  <div class="section">
    <h2>5. Results and Findings</h2>
    <p>{res.get("summary", "Not stated") if isinstance(res, dict) else res}</p>
    {metrics_html}
    {f'<p><strong>Significance:</strong> {res.get("significance")}</p>' if isinstance(res, dict) and res.get("significance") else ""}
  </div>

  <div class="section">
    <h2>6. Limitations</h2>
    <ul>{limits_html}</ul>
  </div>

  <div class="section">
    <h2>7. Future Research Directions</h2>
    <ul>{future_html}</ul>
  </div>

  <div style="margin-top: 40px; text-align: center; font-size: 11px; color: #94a3b8; border-top: 1px solid #f1f5f9; padding-top: 15px;">
    Generated by Research Paper Analyst • Powered by Google Gemini AI
  </div>
</body>
</html>
"""
